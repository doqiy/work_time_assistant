"""Read the supplied field dictionary and timeline sources without inventing fields.

The workbook is read-only. Records contain dictionary fields only; UI metadata is
kept alongside records. A behavior ID is not treated as a unique message ID.
"""

from __future__ import annotations

import json
import math
import os
from collections import OrderedDict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from threading import RLock
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
SOURCE_DIR = BASE_DIR / "data"
FIELDS_PATH = Path(os.getenv("ASSISTANT_FIELDS_PATH", SOURCE_DIR / "字段注释表.xlsx"))
REAL_DATA_PATH = Path(os.getenv("ASSISTANT_REAL_DATA_PATH", SOURCE_DIR / "直播助手频率时长适中主播建议记录_20261020_v2.xlsx"))
HIGH_FREQ_DATA_PATH = Path(os.getenv("ASSISTANT_HIGH_FREQ_DATA_PATH", SOURCE_DIR / "直播助手高频主播建议记录_20261020.xlsx"))
DEMO_PATH = Path(os.getenv("ASSISTANT_DEMO_DATA_PATH", SOURCE_DIR / "demo_data.json"))
ID_FIELDS = {"author_id", "live_stream_id", "user_id", "receiver_author_id", "behavior_id"}
HEADER_ALIASES = {"p_date分区": "p_date", "p_hourmin分区": "p_hourmin"}
SHANGHAI = ZoneInfo("Asia/Shanghai")
EPOCH = datetime(1970, 1, 1, tzinfo=timezone.utc)
_lock = RLock()
_fields_cache: dict[str, Any] = {}
_source_cache: dict[str, Any] = {}


class SourceError(Exception):
    """A required local source is absent or does not follow its declared shape."""


def file_signature(path: Path) -> tuple[str, int, int]:
    try:
        stat = path.stat()
    except OSError as exc:
        raise SourceError(f"无法读取数据源：{path.name}。请检查文件路径和读取权限。") from exc
    return str(path.resolve()), stat.st_mtime_ns, stat.st_size


def read_workbook_rows(path: Path) -> list[tuple]:
    """Reset worksheet dimensions: supplied exports may report only A1."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            worksheet.reset_dimensions()
            return list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    except Exception as exc:
        raise SourceError(f"无法解析工作簿：{path.name}（{type(exc).__name__}）。") from exc


def field_name(raw: Any) -> str:
    name = "" if raw is None else str(raw).strip()
    return HEADER_ALIASES.get(name, name)


def load_fields() -> list[dict[str, str]]:
    with _lock:
        signature = file_signature(FIELDS_PATH)
        if _fields_cache.get("signature") == signature:
            return _fields_cache["fields"]
        rows = read_workbook_rows(FIELDS_PATH)
        if not rows or tuple(rows[0][:4]) != ("字段名", "类型", "字段释义", "备注"):
            raise SourceError("字段注释表表头不符合预期，应为：字段名、类型、字段释义、备注。")
        fields = []
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            values = list(row) + [None] * 4
            fields.append({
                "name": field_name(values[0]),
                "type": str(values[1] or ""),
                "description": str(values[2] or ""),
                "remark": str(values[3] or ""),
            })
        names = [field["name"] for field in fields]
        if not names or len(set(names)) != len(names):
            raise SourceError("字段注释表为空或包含重复字段，无法可靠映射。")
        _fields_cache.update(signature=signature, fields=fields)
        return fields


def normalize_value(name: str, value: Any) -> Any:
    """Keep null, zero and raw enum values; identifiers never pass through float."""
    if value is None:
        return None
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if name in ID_FIELDS:
        if isinstance(value, bool):
            return str(value).lower()
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def project_record(raw: dict[str, Any], allowed: set[str]) -> dict[str, Any]:
    return {name: normalize_value(name, value) for name, value in raw.items() if name in allowed}


def status_key(value: Any) -> str:
    if value is True or (isinstance(value, str) and value.lower() == "true"):
        return "success"
    if value is False or (isinstance(value, str) and value.lower() == "false"):
        return "failure"
    return "unknown"


def numeric_timestamp(value: Any) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(value):
        return value
    return None


def verify_millisecond_timestamps(raw_rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Validate ms against every exported send_time, never by digit length alone.

    send_time is validation-only. It is excluded from timeline records because
    it is not in the supplied field dictionary. This does not establish the
    business meaning of send_timestamp beyond the dictionary's '时间戳'.
    """
    matched = missing = mismatched = 0
    for raw in raw_rows:
        stamp = numeric_timestamp(raw.get("send_timestamp"))
        local_time = raw.get("send_time")
        if stamp is None or not isinstance(local_time, datetime):
            missing += 1
            continue
        if local_time.tzinfo is None:
            local_time = local_time.replace(tzinfo=SHANGHAI)
        try:
            from_stamp = EPOCH + timedelta(milliseconds=stamp)
            if from_stamp == local_time.astimezone(timezone.utc):
                matched += 1
            else:
                mismatched += 1
        except (OverflowError, ValueError):
            mismatched += 1
    return {
        "unit": "ms" if matched and not missing and not mismatched else None,
        "matched": matched,
        "missing": missing,
        "mismatched": mismatched,
        "total": len(raw_rows),
    }


def build_groups(records: list[dict]) -> list[dict]:
    grouped: dict[tuple, dict] = OrderedDict()
    for record in records:
        key = record.get("author_id"), record.get("live_stream_id")
        if key not in grouped:
            grouped[key] = {
                "author_id": key[0], "live_stream_id": key[1], "count": 0,
                "success_count": 0, "failure_count": 0, "unknown_count": 0,
                "earliest": None, "latest": None,
            }
        group = grouped[key]
        group["count"] += 1
        group[status_key(record.get("send_status")) + "_count"] += 1
        stamp = numeric_timestamp(record.get("send_timestamp"))
        if stamp is not None:
            group["earliest"] = stamp if group["earliest"] is None else min(stamp, group["earliest"])
            group["latest"] = stamp if group["latest"] is None else max(stamp, group["latest"])
    return sorted(grouped.values(), key=lambda group: -group["count"])


def _load_real(
    allowed: set[str], path: Path | None = None, source: str = "real",
    display_label: str = "真实数据·频率时长适中",
) -> tuple[list[dict], dict]:
    path = REAL_DATA_PATH if path is None else path
    rows = read_workbook_rows(path)
    if not rows:
        raise SourceError("真实数据工作簿为空。")
    headers = [field_name(value) for value in rows[0]]
    if len(headers) != len(set(headers)):
        raise SourceError("真实数据存在重复表头，无法可靠映射。")
    required = {"author_id", "live_stream_id", "send_timestamp"}
    if required - set(headers):
        raise SourceError("真实数据缺少分组或排序字段：" + "、".join(sorted(required - set(headers))))
    raw_rows = [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]
    validation = verify_millisecond_timestamps(raw_rows)
    records = [project_record(raw, allowed) for raw in raw_rows]
    notes = [
        "仅使用《字段注释表》中的字段；send_time 只用于校验时间戳，from_scene 不参与展示或解释。",
        "send_timestamp 的字段释义为“时间戳”；不将它断言为行为发生时刻。",
        "发送成功不等于主播已看到；按原始行保留记录，不按 behavior_id 或相同文案去重。",
    ]
    if validation["unit"] == "ms":
        notes.append(f"全表 {validation['matched']:,} 条 send_timestamp 按毫秒转为北京时间，与导出 send_time 逐条一致。")
    else:
        notes.append(f"时间戳单位未通过全表验证：一致 {validation['matched']} 条，缺失 {validation['missing']} 条，不一致 {validation['mismatched']} 条。保留原值，暂停时间换算。")
    if "user_id" not in headers:
        notes.append("本文件缺少 user_id，用户显示“未提供”，不从行为 ID 或文案反推用户。")
    if records and all(status_key(row.get("send_status")) == "success" for row in records):
        notes.append("本文件所有记录均为发送成功；失败和未知状态的展示可在模拟数据中查看。")
    enum_mixed = [name for name in ("behavior_type", "is_topic") if len({type(row.get(name)) for row in records if row.get(name) is not None}) > 1]
    if enum_mixed:
        notes.append("以下字段同时含中文文本与数值码，保留原值，不反向补码：" + "、".join(enum_mixed) + "。")
    large_numeric_ids = sum(1 for row in raw_rows for name in ID_FIELDS if isinstance(row.get(name), (int, float)) and not isinstance(row.get(name), bool) and abs(row[name]) >= 10**15)
    if large_numeric_ids:
        notes.append("部分 ID 在 Excel 中以长数字保存，源文件可能已损失精度；本页保留读取值并转文本，无法恢复导出前的原 ID。")
    user_ids = [row["user_id"] for row in records if row.get("user_id") not in (None, "")]
    return records, {
        "source": source, "title": path.name, "display_label": display_label, "is_demo": False,
        "record_count": len(records), "missing_fields": sorted(allowed - set(headers)),
        "user_id_record_count": len(user_ids), "unique_user_count": len(set(user_ids)),
        "extra_fields": sorted(set(headers) - allowed), "notes": notes,
        "timestamp_unit": validation["unit"], "timestamp_validation": validation,
        "timezone": "Asia/Shanghai", "field_count": len(set(headers) & allowed),
    }


def _load_demo(allowed: set[str]) -> tuple[list[dict], dict]:
    try:
        data = json.loads(DEMO_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise SourceError("无法读取模拟数据 demo_data.json。") from exc
    raw_records = data.get("records")
    if not isinstance(raw_records, list) or not all(isinstance(row, dict) for row in raw_records):
        raise SourceError("模拟数据 records 必须为记录数组。")
    provided = {key for record in raw_records for key in record}
    if provided - allowed:
        raise SourceError("模拟数据出现字典外字段：" + "、".join(sorted(provided - allowed)))
    records = [project_record(raw, allowed) for raw in raw_records]
    return records, {
        "source": "demo", "title": "模拟数据 · 仅用于展示效果", "display_label": "模拟数据", "is_demo": True,
        "record_count": len(records), "missing_fields": sorted(allowed - provided),
        "extra_fields": [], "timestamp_unit": "ms", "timezone": "Asia/Shanghai",
        "field_count": len(provided),
        "notes": [
            "全部用户、主播、直播间和提示内容均为模拟，不代表真实触发或效果。",
            "模拟数据明确采用 Unix 毫秒时间戳；主时间轴展示北京时间。",
            "记录严格使用字段注释表中的字段；发送成功不等于已曝光。",
            "包含同一时间多条、长间隔、发送失败、状态未知和嘉宾端，用于检查展示。",
        ],
    }


def load_source(source: str = "demo") -> dict:
    if source not in {"demo", "real", "real_high"}:
        raise SourceError("不支持的数据源，请选择 demo、real 或 real_high。")
    with _lock:
        fields = load_fields()
        path = {"demo": DEMO_PATH, "real": REAL_DATA_PATH, "real_high": HIGH_FREQ_DATA_PATH}[source]
        signature = (file_signature(path), file_signature(FIELDS_PATH))
        cached = _source_cache.get(source)
        if cached and cached["signature"] == signature:
            return cached["dataset"]
        allowed = {field["name"] for field in fields}
        if source == "demo":
            records, meta = _load_demo(allowed)
        else:
            label = "真实数据·高频主播" if source == "real_high" else "真实数据·频率时长适中"
            records, meta = _load_real(allowed, path, source, label)
        # Python's sort is stable: equal timestamps keep source order, no dedup.
        records.sort(key=lambda row: (
            numeric_timestamp(row.get("send_timestamp")) is None,
            numeric_timestamp(row.get("send_timestamp")) or 0,
        ))
        dataset = {"meta": meta, "records": records, "groups": build_groups(records)}
        _source_cache[source] = {"signature": signature, "dataset": dataset}
        return dataset


def source_summary(source: str = "demo") -> dict:
    data = load_source(source)
    return {"meta": data["meta"], "groups": data["groups"]}


def source_records(author_id: str, live_stream_id: str | None = None, source: str = "demo") -> dict:
    data = load_source(source)
    records = [row for row in data["records"] if row.get("author_id") == author_id
               and (live_stream_id is None or row.get("live_stream_id") == live_stream_id)]
    return {"meta": data["meta"], "records": records}
